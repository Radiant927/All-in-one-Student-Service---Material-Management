from fastapi import APIRouter, Depends, Query
from sqlalchemy.orm import Session
from database import get_db
from services.report_service import (
    generate_restock_report,
    get_consumption_trend,
    get_dashboard_alerts,
    predict_stockout,
)
from services.email_service import send_restock_email
import io
from openpyxl import Workbook
from fastapi.responses import StreamingResponse
from datetime import datetime

router = APIRouter()


@router.get("/reports/restock-suggestions")
def restock_suggestions(
    window_days: int = Query(30, ge=7, le=90),
    db: Session = Depends(get_db)
):
    """Generate restock suggestion report."""
    report = generate_restock_report(db, window_days)
    return {"ok": True, "data": report, "msg": ""}


@router.get("/reports/consumption-trends")
def consumption_trends(
    material_id: int = Query(...),
    months: int = Query(3, ge=1, le=12),
    db: Session = Depends(get_db)
):
    """Get monthly consumption trend for a material."""
    trend = get_consumption_trend(db, material_id, months)
    if trend is None:
        return {"ok": False, "data": None, "msg": "物资不存在"}
    return {"ok": True, "data": trend, "msg": ""}


@router.get("/reports/dashboard-alerts")
def dashboard_alerts(
    window_days: int = Query(30, ge=7, le=90),
    db: Session = Depends(get_db)
):
    """Get dashboard alert summary."""
    alerts = get_dashboard_alerts(db, window_days)
    return {"ok": True, "data": alerts, "msg": ""}


@router.get("/reports/material-alert/{material_id}")
def material_alert(
    material_id: int,
    window_days: int = Query(30, ge=7, le=90),
    db: Session = Depends(get_db)
):
    """Get detailed alert/prediction for a single material."""
    pred = predict_stockout(db, material_id, window_days)
    if pred is None:
        return {"ok": False, "data": None, "msg": "物资不存在"}
    return {"ok": True, "data": pred, "msg": ""}


@router.get("/reports/export")
def export_report(
    window_days: int = Query(30, ge=7, le=90),
    format: str = Query("excel"),
    db: Session = Depends(get_db)
):
    """Export restock report as Excel file."""
    if format != "excel":
        return {"ok": False, "data": None, "msg": "不支持的格式"}

    report = generate_restock_report(db, window_days)

    wb = Workbook()
    ws = wb.active
    ws.title = "补货建议报表"

    # Header row
    headers = ["物资名称", "规格", "单位", "当前库存", "日均消耗", "预计耗尽(天)", "建议补货量", "紧急程度", "数据置信度"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    # Data rows
    urgency_labels = {"critical": "紧急", "warning": "预警", "ok": "正常", "insufficient_data": "数据不足"}
    for row_idx, item in enumerate(report["items"], 2):
        ws.cell(row=row_idx, column=1, value=f"{item['icon']} {item['name']}")
        ws.cell(row=row_idx, column=2, value=item["spec"])
        ws.cell(row=row_idx, column=3, value=item["unit"])
        ws.cell(row=row_idx, column=4, value=item["current_stock"])
        ws.cell(row=row_idx, column=5, value=item["daily_velocity"])
        ws.cell(row=row_idx, column=6, value=item["days_until_empty"] if item["days_until_empty"] is not None else "N/A")
        ws.cell(row=row_idx, column=7, value=item["suggested_restock_qty"])
        ws.cell(row=row_idx, column=8, value=urgency_labels.get(item["status"], item["status"]))
        ws.cell(row=row_idx, column=9, value=item["data_confidence"])

    # Auto-width (approximate)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=restock_report.xlsx"}
    )


_LAST_EMAIL_SENT = None


@router.post("/reports/send-email")
def send_email(
    window_days: int = Query(30, ge=7, le=90),
    db: Session = Depends(get_db)
):
    """Send restock report via email."""
    global _LAST_EMAIL_SENT
    result = send_restock_email(db, window_days=window_days)
    if result["ok"]:
        _LAST_EMAIL_SENT = datetime.now().isoformat()
    return result


@router.get("/reports/email-status")
def email_status():
    """Get the last email send timestamp."""
    return {
        "ok": True,
        "data": {"last_sent_at": _LAST_EMAIL_SENT},
        "msg": ""
    }
