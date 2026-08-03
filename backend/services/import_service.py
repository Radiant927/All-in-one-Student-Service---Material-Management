from io import BytesIO
from datetime import datetime
from sqlalchemy.orm import Session
from models import Material, Warehouse, StorageLocation, InventoryBatch, InventoryItem


def import_from_excel(db: Session, file_bytes: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(file_bytes))
    ws = wb.active

    # Read headers
    headers = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value:
            headers[cell.value.strip()] = col_idx

    result = {"imported": 0, "skipped": 0, "errors": []}
    now = datetime.now().isoformat()

    # Get or create warehouses
    default_wh = db.query(Warehouse).first()
    if not default_wh:
        default_wh = Warehouse(name="主仓库", location_desc="一楼库房")
        db.add(default_wh)
        db.flush()

    for row_idx in range(2, ws.max_row + 1):
        try:
            row = {h: ws.cell(row=row_idx, column=col).value for h, col in headers.items()}
            name = row.get("品名", row.get("name", ""))
            if not name or not str(name).strip():
                continue

            name = str(name).strip()
            spec = str(row.get("规格", row.get("spec", ""))).strip()
            unit = str(row.get("单位", row.get("unit", "个"))).strip() or "个"
            category = str(row.get("分类", row.get("category", "consumable"))).strip()
            sub_category = str(row.get("子分类", row.get("sub_category", "direct_consumption"))).strip()
            has_tracking = str(row.get("个体追踪", row.get("has_individual_tracking", "否"))).strip()
            quantity = int(row.get("数量", row.get("quantity", 0)) or 0)
            warehouse_name = str(row.get("仓库", row.get("warehouse", ""))).strip()
            location_code = str(row.get("储位", row.get("location", ""))).strip()
            threshold = int(row.get("预警阈值", row.get("low_stock_threshold", 5)) or 5)

            # Map category
            if category in ("固定", "固定性物资", "durable"):
                category = "durable"
            else:
                category = "consumable"

            # Map sub_category
            sub_map = {
                "全新消耗品": "new_consumable", "可循环物资": "recyclable",
                "直接消耗品": "direct_consumption", "new_consumable": "new_consumable",
                "recyclable": "recyclable", "direct_consumption": "direct_consumption"
            }
            sub_category = sub_map.get(sub_category, "direct_consumption")

            has_tracking = has_tracking in ("是", "true", "True", "1", "yes")

            # Find or create warehouse
            wh = default_wh
            if warehouse_name:
                wh = db.query(Warehouse).filter(Warehouse.name == warehouse_name).first()
                if not wh:
                    wh = Warehouse(name=warehouse_name, location_desc="")
                    db.add(wh)
                    db.flush()

            # Find or create location
            loc = None
            if location_code:
                loc = db.query(StorageLocation).filter(
                    StorageLocation.warehouse_id == wh.id,
                    StorageLocation.full_code == location_code
                ).first()
                if not loc:
                    parts = location_code.split("-")
                    shelf = parts[0] if len(parts) > 0 else "A"
                    level = int(parts[1]) if len(parts) > 1 else 1
                    position = parts[2] if len(parts) > 2 else ""
                    loc = StorageLocation(warehouse_id=wh.id, shelf=shelf, level=level,
                                          position=position, full_code=location_code)
                    db.add(loc)
                    db.flush()

            # Check if material already exists
            existing = db.query(Material).filter(Material.name == name).first()
            if existing:
                # Update inventory
                if existing.has_individual_tracking:
                    if not has_tracking:
                        result["skipped"] += 1
                        continue
                else:
                    if quantity > 0:
                        batch = db.query(InventoryBatch).filter(
                            InventoryBatch.material_id == existing.id,
                            InventoryBatch.warehouse_id == wh.id
                        ).first()
                        if batch:
                            batch.quantity += quantity
                            batch.updated_at = now
                        else:
                            db.add(InventoryBatch(material_id=existing.id, warehouse_id=wh.id,
                                                  location_id=loc.id if loc else None,
                                                  quantity=quantity, created_at=now, updated_at=now))
                result["imported"] += 1
                continue

            # Create material
            m = Material(
                name=name, spec=spec, unit=unit, category=category,
                sub_category=sub_category,
                has_individual_tracking=1 if has_tracking else 0,
                low_stock_threshold=threshold, icon="📦", color_idx=0,
                created_at=now, updated_at=now
            )
            db.add(m)
            db.flush()

            # Create inventory
            if has_tracking and quantity > 0:
                pad_len = max(3, len(str(quantity)))
                for i in range(1, quantity + 1):
                    code = f"{name[:2].upper()}-{str(i).zfill(pad_len)}"
                    db.add(InventoryItem(material_id=m.id, code=code, status="available",
                                         warehouse_id=wh.id, location_id=loc.id if loc else None,
                                         created_at=now))
            elif quantity > 0:
                db.add(InventoryBatch(material_id=m.id, warehouse_id=wh.id,
                                      location_id=loc.id if loc else None,
                                      quantity=quantity, created_at=now, updated_at=now))

            result["imported"] += 1
        except Exception as e:
            result["errors"].append(f"第{row_idx}行: {str(e)}")
            result["skipped"] += 1
            continue

    db.commit()
    return result